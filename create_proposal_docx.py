# -*- coding: utf-8 -*-
"""
Script to generate the official Proposal Document for Kopi Jodi Digital Ecosystem.
Outputs: Penawaran_Pengembangan_Ekosistem_Aplikasi_Kopi_Jodi.docx
Author: genossys (genossys2019@gmail.com)
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Set background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag.endswith('shd'):
            tcPr.remove(child)
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shading)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    """Set padding in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'  <w:top w:w="{top}" w:type="dxa"/>'
        f'  <w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'  <w:left w:w="{left}" w:type="dxa"/>'
        f'  <w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_cell_borders(cell, top="E2D9CC", bottom="E2D9CC", left="none", right="none", 
                     sz="4", left_sz="24"):
    """Set borders on cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag.endswith('tcBorders'):
            tcPr.remove(child)
    
    left_xml = f'<w:left w:val="single" w:sz="{left_sz}" w:space="0" w:color="{left}"/>' if left != "none" else '<w:left w:val="none"/>'
    right_xml = f'<w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{right}"/>' if right != "none" else '<w:right w:val="none"/>'
    top_xml = f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{top}"/>' if top != "none" else '<w:top w:val="none"/>'
    bottom_xml = f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{bottom}"/>' if bottom != "none" else '<w:bottom w:val="none"/>'
    
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  {top_xml}'
        f'  {left_xml}'
        f'  {bottom_xml}'
        f'  {right_xml}'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

def add_callout(doc, title, text_lines, border_color="C27835", bg_color="FAF7F2"):
    """Create a callout box with a colored left accent border."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.27)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=120, bottom=120, left=160, right=140)
    set_cell_borders(cell, top="none", bottom="none", left=border_color, right="none", sz="0", left_sz="24")
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    run_title = p.add_run(title)
    run_title.bold = True
    run_title.font.name = 'Segoe UI'
    run_title.font.size = Pt(10)
    run_title.font.color.rgb = RGBColor(0x3E, 0x27, 0x23)
    
    for line in text_lines:
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_before = Pt(1)
        p2.paragraph_format.space_after = Pt(2)
        p2.paragraph_format.line_spacing = 1.15
        run = p2.add_run(line)
        run.font.name = 'Segoe UI'
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0x37, 0x41, 0x51)
        
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def format_table(table, col_widths, col_alignments, header_bg="2C1810", zebra=True):
    """Format table with headers, zebra striping, padding, and subtle borders."""
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    header_row = table.rows[0]
    trPr = header_row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    
    for idx, cell in enumerate(header_row.cells):
        cell.width = Inches(col_widths[idx])
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_background(cell, header_bg)
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        set_cell_borders(cell, top="2C1810", bottom="C27835", left="none", right="none", sz="6")
        for p in cell.paragraphs:
            p.alignment = col_alignments[idx]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            for run in p.runs:
                run.bold = True
                run.font.name = 'Segoe UI'
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                
    for r_idx, row in enumerate(table.rows[1:], start=1):
        bg = "FBF9F5" if (zebra and r_idx % 2 == 1) else "FFFFFF"
        for c_idx, cell in enumerate(row.cells):
            cell.width = Inches(col_widths[c_idx])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=90, bottom=90, left=120, right=120)
            set_cell_borders(cell, top="E5E0D8", bottom="E5E0D8", left="none", right="none", sz="4")
            for p in cell.paragraphs:
                p.alignment = col_alignments[c_idx]
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.15
                for run in p.runs:
                    run.font.name = 'Segoe UI'
                    run.font.size = Pt(9)
                    if not run.font.color.rgb:
                        run.font.color.rgb = RGBColor(0x37, 0x41, 0x51)

def style_heading(p, text, level=1):
    """Custom styled heading with elegant coffee accent."""
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    run.font.name = 'Segoe UI'
    
    if level == 1:
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(5)
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(0x2C, 0x18, 0x10) # Deep Espresso
        pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="10" w:space="3" w:color="C27835"/></w:pBdr>')
        p._p.get_or_add_pPr().append(pBdr)
    elif level == 2:
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        run.font.size = Pt(11.5)
        run.font.color.rgb = RGBColor(0x4A, 0x2E, 0x18) # Warm Mocha
    elif level == 3:
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0xC2, 0x78, 0x35) # Warm Caramel
    return run

def add_body_p(doc, text="", bold_prefix="", space_after=3):
    """Add a clean body paragraph."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.name = 'Segoe UI'
        r_pre.font.size = Pt(9.5)
        r_pre.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
        
    if text:
        r_text = p.add_run(text)
        r_text.font.name = 'Segoe UI'
        r_text.font.size = Pt(9.5)
        r_text.font.color.rgb = RGBColor(0x37, 0x41, 0x51)
        
    return p

def add_bullet_p(doc, bold_prefix, text, space_after=2):
    """Add a bullet point."""
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.name = 'Segoe UI'
        r_pre.font.size = Pt(9)
        r_pre.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
        
    r_text = p.add_run(text)
    r_text.font.name = 'Segoe UI'
    r_text.font.size = Pt(9)
    r_text.font.color.rgb = RGBColor(0x37, 0x41, 0x51)
    return p

def build_proposal_docx(file_path):
    doc = docx.Document()
    
    # 1. Page Setup (A4)
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Header setup
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("KOPI JODI x GENOSSYS | Proposal Ringkas Ekosistem Aplikasi Terpadu")
        hrun.font.name = 'Segoe UI'
        hrun.font.size = Pt(8)
        hrun.font.color.rgb = RGBColor(0x9C, 0xA3, 0xAF)
        
        # Footer setup
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        frun1 = fp.add_run("Dipersiapkan oleh genossys (genossys2019@gmail.com) | Khusus Manajemen Kopi Jodi")
        frun1.font.name = 'Segoe UI'
        frun1.font.size = Pt(8)
        frun1.font.color.rgb = RGBColor(0x9C, 0xA3, 0xAF)

    # 2. Cover / Title Box with Genossys Branding
    title_table = doc.add_table(rows=1, cols=1)
    title_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_cell = title_table.cell(0, 0)
    t_cell.width = Inches(6.27)
    set_cell_background(t_cell, "2C1810")
    set_cell_margins(t_cell, top=200, bottom=220, left=260, right=260)
    set_cell_borders(t_cell, top="C27835", bottom="C27835", left="C27835", right="C27835", sz="10")
    
    # Genossys Logo in Header Box
    logo_path = r"c:\PROJECT\WEBSITE\kopi-jodi\genossys_horizontal_dark.png"
    if os.path.exists(logo_path):
        tp_logo = t_cell.paragraphs[0]
        tp_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        tp_logo.paragraph_format.space_before = Pt(0)
        tp_logo.paragraph_format.space_after = Pt(4)
        run_logo = tp_logo.add_run()
        run_logo.add_picture(logo_path, width=Inches(2.5))
        tp0 = t_cell.add_paragraph()
    else:
        tp0 = t_cell.paragraphs[0]

    tp0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tp0.paragraph_format.space_before = Pt(2)
    tp0.paragraph_format.space_after = Pt(2)
    tr0 = tp0.add_run("PROPOSAL PENAWARAN TEKNOLOGI & APLIKASI")
    tr0.bold = True
    tr0.font.name = 'Segoe UI'
    tr0.font.size = Pt(9.5)
    tr0.font.color.rgb = RGBColor(0xD9, 0x9B, 0x62)
    
    tp1 = t_cell.add_paragraph()
    tp1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tp1.paragraph_format.space_before = Pt(4)
    tp1.paragraph_format.space_after = Pt(4)
    tr1 = tp1.add_run("EKOSISTEM DIGITAL KOPI JODI")
    tr1.bold = True
    tr1.font.name = 'Segoe UI'
    tr1.font.size = Pt(19)
    tr1.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    
    tp2 = t_cell.add_paragraph()
    tp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tp2.paragraph_format.space_before = Pt(2)
    tp2.paragraph_format.space_after = Pt(6)
    tr2 = tp2.add_run("Sistem Terpadu: ERP Backoffice, POS Kasir, Kitchen Display (KDS), Mobile App, & Portal Kemitraan")
    tr2.font.name = 'Segoe UI'
    tr2.font.size = Pt(9.5)
    tr2.font.color.rgb = RGBColor(0xEA, 0xE5, 0xDD)
    
    tp3 = t_cell.add_paragraph()
    tp3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tp3.paragraph_format.space_before = Pt(2)
    tr3 = tp3.add_run("Pilihan Skema: Bulanan Dedicated Dev (Rp 5 Jt/bln, Min. 5 Thn) & Beli Putus (Turnkey)")
    tr3.bold = True
    tr3.font.name = 'Segoe UI'
    tr3.font.size = Pt(9)
    tr3.font.color.rgb = RGBColor(0xF5, 0x9E, 0x0B)
    
    # Metadata info
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    meta_table = doc.add_table(rows=2, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    col_w = [3.13, 3.14]
    
    # Left: Client
    p_cl = meta_table.rows[0].cells[0].paragraphs[0]
    p_cl.add_run("Dipersiapkan Untuk: ").bold = True
    p_cl.add_run("Founder & Manajemen Kopi Jodi")
    
    # Right: Developer (genossys)
    p_dev = meta_table.rows[0].cells[1].paragraphs[0]
    p_dev.add_run("Disusun Oleh: ").bold = True
    p_dev.add_run("genossys (Technology Partner)")
    
    # Left bottom: Date & Format
    p_dt = meta_table.rows[1].cells[0].paragraphs[0]
    p_dt.add_run("Versi & Tanggal: ").bold = True
    p_dt.add_run("Oktober 2026 | Versi Ringkas")
    
    # Right bottom: Email
    p_em = meta_table.rows[1].cells[1].paragraphs[0]
    p_em.add_run("Email Kontak: ").bold = True
    p_em.add_run("genossys2019@gmail.com")
    
    for r in meta_table.rows:
        for c_idx, c in enumerate(r.cells):
            c.width = Inches(col_w[c_idx])
            set_cell_background(c, "FDFBF7")
            set_cell_margins(c, top=50, bottom=50, left=70, right=70)
            set_cell_borders(c, top="E2D9CC", bottom="E2D9CC", left="E2D9CC", right="E2D9CC", sz="4")
            p = c.paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            for r_run in p.runs:
                r_run.font.name = 'Segoe UI'
                r_run.font.size = Pt(8.5)
                
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # -------------------------------------------------------------
    # SECTION 1: MASALAH OPERASIONAL & SOLUSI
    # -------------------------------------------------------------
    style_heading(doc.add_paragraph(), "1. Masalah Operasional & Solusi Terpadu", level=1)
    
    add_bullet_p(doc, "Kebocoran Stok Bahan Baku: ", "Pengurangan fisik biji kopi, susu, sirup, dan cup sering tidak cocok dengan transaksi kasir.")
    add_bullet_p(doc, "HPP Bias & Lambat: ", "Nota belanja fisik darurat outlet sering terlambat atau tidak diverifikasi Finance pusat.")
    add_bullet_p(doc, "Risiko Fraud Kas Kecil: ", "Pengeluaran darurat outlet belum memiliki batas plafon approval berjenjang.")
    add_bullet_p(doc, "Laporan Cabang Lambat: ", "Rekapitulasi omzet dan perhitungan bagi hasil dengan mitra/investor masih manual.")
    
    add_callout(doc, "Solusi Ekosistem Kopi Jodi oleh genossys:", [
        "Membangun 1 platform terintegrasi (seperti Fore / Kopi Kenangan) yang menghubungkan Kasir POS, Layar Barista (KDS), Resep Otomatis, Approval Pembelian & Kas Kecil, hingga Portal Investor dalam satu database real-time."
    ], border_color="C27835", bg_color="FAF7F2")

    # -------------------------------------------------------------
    # SECTION 2: CAKUPAN MODUL & APLIKASI
    # -------------------------------------------------------------
    style_heading(doc.add_paragraph(), "2. Cakupan Modul & Aplikasi", level=1)

    app_table = doc.add_table(rows=9, cols=5)
    headers = ["#", "Modul Aplikasi", "Platform & Pengguna", "Fungsi Utama", "Fase"]
    col_widths = [0.35, 1.4, 1.35, 2.6, 0.57]
    col_align = [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER]
    
    for i, h in enumerate(headers):
        app_table.cell(0, i).paragraphs[0].text = h
        
    app_data = [
        ("1", "ERP Backoffice", "Web Responsive\n(Owner, Finance, Admin)", "Master menu & resep, stok otomatis, verifikasi nota, kas kecil bertingkat, & laba/rugi cabang.", "Fase 1"),
        ("2", "POS Kasir", "Tablet / PC Web\n(Kasir Outlet)", "Transaksi cepat, modifier (ice/sugar/size), kas laci, mode offline, cetak struk, QRIS Midtrans.", "Fase 1"),
        ("3", "Layar Barista (KDS)", "Tablet / Monitor\n(Barista)", "Antrean pesanan real-time, status pesanan (Queue -> In Progress -> Ready), tanpa kertas bon.", "Fase 1"),
        ("4", "App Operasi Outlet", "Mobile Android\n(Store Manager)", "Penerimaan barang (surat jalan), stock opname fisik, pencatatan waste, pengajuan kas kecil.", "Fase 1"),
        ("5", "App Pelanggan", "iOS & Android\n(Pelanggan Kopi Jodi)", "Order online (Pick-up / Delivery), QRIS/E-Wallet, loyalty poin, e-voucher promo.", "Fase 2"),
        ("6", "CRM & Promo Panel", "Web Backoffice\n(Tim Marketing)", "Kelola voucher diskon, blast push notification penawaran promo, segmentasi pelanggan.", "Fase 2"),
        ("7", "Portal Mitra / Investor", "Web Portal (Secure)\n(Mitra Pemilik Cabang)", "Laporan real-time: omzet cabang, rincian biaya, HPP riil, dan estimasi bagi hasil.", "Fase 3"),
        ("8", "Owner Dashboard & HR", "Web / Mobile\n(Owner & HR)", "KPI konsolidasi seluruh cabang, jadwal shift barista, & absensi GPS karyawan.", "Fase 3")
    ]
    
    for row_idx, data in enumerate(app_data, start=1):
        row = app_table.rows[row_idx]
        for col_idx, text in enumerate(data):
            row.cells[col_idx].paragraphs[0].text = text
            
    format_table(app_table, col_widths, col_align, header_bg="2C1810")
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Fitur Kunci
    style_heading(doc.add_paragraph(), "Fitur Kunci Pengendalian Biaya & Stok (ERP Core):", level=2)
    add_bullet_p(doc, "Potong Stok Resep Otomatis: ", "Tiap transaksi kasir langsung memotong gramatur kopi, ml susu, sirup, dan cup secara otomatis.")
    add_bullet_p(doc, "Alur Pembelian Transparan (True-Up HPP): ", "Request bahan (tanpa harga) -> Terima barang fisik -> Upload foto bon -> Finance verifikasi -> HPP riil terkoreksi otomatis.")
    add_bullet_p(doc, "Approval Kas Kecil Berjenjang: ", "Nominal kecil (Store Manager), nominal menengah (Finance), nominal besar (Owner) wajib lampirkan foto bon.")
    add_bullet_p(doc, "Multi-Outlet & Multi-Investor: ", "Pemisahan pembukuan cabang milik sendiri vs mitra waralaba secara transparan.")

    # -------------------------------------------------------------
    # SECTION 3: SKEMA KERJASAMA & BIAYA
    # -------------------------------------------------------------
    style_heading(doc.add_paragraph(), "3. Pilihan Skema Kerjasama & Biaya", level=1)
    add_body_p(doc, "Kami menyediakan 2 alternatif skema yang dapat disesuaikan dengan kebutuhan dan strategi cashflow Kopi Jodi:")

    # OPSI 1: Beli Putus
    style_heading(doc.add_paragraph(), "OPSI 1: Beli Putus — Turnkey Fixed-Price (Kepastian Budget Sekali Bayar)", level=2)
    add_body_p(doc, "Proyek tuntas dengan anggaran terkunci di awal, penyerahan penuh source code, dan garansi bebas bug setelah serah terima:")

    bp_table = doc.add_table(rows=5, cols=4)
    bp_headers = ["Paket", "Cakupan Modul Aplikasi", "Estimasi Waktu", "Biaya Standar"]
    bp_widths = [1.2, 2.77, 1.1, 1.2]
    bp_align = [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT]
    
    for i, h in enumerate(bp_headers):
        bp_table.cell(0, i).paragraphs[0].text = h
        
    bp_data = [
        ("Fase 1 (Core)", "ERP Backoffice, POS Kasir, KDS Barista, App Operasi Outlet", "2.5 - 3 Bulan", "Rp 75.000.000,-"),
        ("Fase 2 (Omni)", "Mobile App Pelanggan (iOS & Android) + Panel Promo CRM", "2 Bulan", "Rp 48.000.000,-"),
        ("Fase 3 (Scale)", "Portal Mitra/Investor + Owner Dashboard & HR Shift", "1 - 1.5 Bulan", "Rp 27.000.000,-"),
        ("PAKET BUNDLING\n(Fase 1 + 2 + 3)", "Seluruh Ekosistem Lengkap (Diskon Khusus Rp 15 Juta)", "± 6 Bulan", "Rp 135.000.000,-\n(Net)")
    ]
    
    for r_idx, d in enumerate(bp_data, start=1):
        row = bp_table.rows[r_idx]
        for c_idx, val in enumerate(d):
            row.cells[c_idx].paragraphs[0].text = val
            
    # Highlight bundling row
    for c in bp_table.rows[4].cells:
        for p in c.paragraphs:
            for r in p.runs:
                r.bold = True
                
    format_table(bp_table, bp_widths, bp_align, header_bg="2C1810")
    
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    add_body_p(doc, "*Fasilitas Beli Putus: Garansi bug 3-6 bulan, penyerahan full source code, setup cloud, SOP, dan pelatihan staf.", space_after=6)

    # OPSI 2: Bulanan
    add_callout(doc, "OPSI 2: Sistem Bulanan — Dedicated Programmer (Disarankan untuk Cashflow Ringan & Bebas Tambah Fitur)", [
        "• Biaya Bulanan: Rp 5.000.000,- / bulan (Flat)",
        "• Komitmen Kontrak: Minimal 5 Tahun (60 Bulan)",
        "• Total Nilai Kontrak (5 Tahun): Rp 300.000.000,- (dicicil flat Rp 5 Juta per bulan)",
        "• BEBAS TAMBAH FITUR APAPUN (Unlimited Custom Features): Kopi Jodi bebas meminta penambahan modul/fitur baru atau integrasi kustom kapan saja tanpa batasan dan tanpa biaya tambahan (Zero Change Request Fee).",
        "• Tim Dedicated dari genossys: Membangun seluruh ekosistem aplikasi secara bertahap (Fase 1, 2, hingga 3).",
        "• Maintenance & Support Siaga 5 Tahun: Pemeliharaan rutin, bug fixing, dan server monitoring standby selama 5 tahun.",
        "• Bebas Biaya HR: Tanpa beban rekrutmen, THR, BPJS, pesangon, maupun penyediaan laptop kerja programmer.",
        "• Hak Milik: Seluruh source code yang dikembangkan menjadi milik Kopi Jodi."
    ], border_color="059669", bg_color="ECFDF5")

    # Tabel Perbandingan
    style_heading(doc.add_paragraph(), "Perbandingan Cepat: Opsi 1 vs Opsi 2", level=2)
    
    cmp_table = doc.add_table(rows=6, cols=3)
    cmp_headers = ["Parameter", "Opsi 1: Beli Putus (Turnkey)", "Opsi 2: Bulanan (Dedicated Programmer)"]
    cmp_widths = [1.7, 2.3, 2.27]
    cmp_align = [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT]
    
    for i, h in enumerate(cmp_headers):
        cmp_table.cell(0, i).paragraphs[0].text = h
        
    cmp_data = [
        ("Beban Cashflow", "Pengeluaran modal per termin proyek", "Sangat Ringan (Rp 5 Juta/bln flat)"),
        ("Masa Komitmen", "Selesai per milestone fase (~3 - 6 bulan)", "Minimal 5 Tahun (Kerjasama Jangka Panjang)"),
        ("Maintenance & Update", "Garansi 3 - 6 bulan (setelahnya kontrak terpisah)", "Gratis & Standby selama 5 Tahun oleh genossys"),
        ("Fleksibilitas Fitur", "Terkunci sesuai scope dokumen spesifikasi awal", "100% Bebas Tambah Fitur Apapun Kapan Saja (tanpa biaya ekstra per request)"),
        ("Rekomendasi", "Sangat Tepat jika Kopi Jodi memiliki modal awal & ingin proyek selesai dalam target waktu.", "Sangat Tepat jika Kopi Jodi ingin inovasi bebas tanpa batas & cashflow bulanan ringan.")
    ]
    
    for r_idx, d in enumerate(cmp_data, start=1):
        row = cmp_table.rows[r_idx]
        for c_idx, val in enumerate(d):
            row.cells[c_idx].paragraphs[0].text = val
            
    format_table(cmp_table, cmp_widths, cmp_align, header_bg="2C1810")
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # -------------------------------------------------------------
    # SECTION 4: ESTIMASI BIAYA PIHAK KETIGA
    # -------------------------------------------------------------
    style_heading(doc.add_paragraph(), "4. Estimasi Biaya Layanan Pihak Ketiga (Infrastruktur)", level=1)
    add_body_p(doc, "Biaya operasional pihak ketiga dibayarkan langsung ke penyedia resmi (tanpa markup):")

    tp_table = doc.add_table(rows=6, cols=4)
    tp_headers = ["Layanan", "Penyedia", "Estimasi Biaya", "Keterangan Siklus"]
    tp_widths = [1.8, 1.4, 1.5, 1.57]
    tp_align = [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.LEFT]
    
    for i, h in enumerate(tp_headers):
        tp_table.cell(0, i).paragraphs[0].text = h
        
    tp_data = [
        ("Cloud Server VPS & Database", "DigitalOcean / AWS Lightsail", "Rp 600.000 - Rp 1.200.000 / bln", "Bulanan (sesuai traffic outlet)"),
        ("Payment Gateway (QRIS)", "Midtrans Indonesia", "0.7% per transaksi QRIS", "Dipotong per transaksi (tanpa biaya bulanan)"),
        ("Google Play Developer Account", "Google LLC", "$25 USD (~Rp 400.000)", "Sekali bayar seumur hidup (App Android)"),
        ("Apple Developer Program", "Apple Inc.", "$99 USD (~Rp 1.600.000) / thn", "Tahunan (App iOS iPhone)"),
        ("Domain Web & Sertifikat SSL", "Cloudflare / Niagahoster", "~Rp 250.000 / thn", "Tahunan (SSL otomatis gratis)")
    ]
    
    for r_idx, d in enumerate(tp_data, start=1):
        row = tp_table.rows[r_idx]
        for c_idx, val in enumerate(d):
            row.cells[c_idx].paragraphs[0].text = val
            
    format_table(tp_table, tp_widths, tp_align, header_bg="2C1810")
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # -------------------------------------------------------------
    # SECTION 5: TIMELINE & KETENTUAN
    # -------------------------------------------------------------
    style_heading(doc.add_paragraph(), "5. Timeline Pengerjaan & Ketentuan Pembayaran", level=1)
    
    add_bullet_p(doc, "Minggu 1 - 2: ", "Finalisasi alur operasional outlet, master resep, dan perancangan database.")
    add_bullet_p(doc, "Minggu 3 - 8: ", "Pengembangan Core ERP Backoffice, POS Tablet kasir, dan Layar Barista (KDS).")
    add_bullet_p(doc, "Minggu 9 - 10: ", "Integrasi hardware kasir (printer thermal, QRIS) dan Quality Assurance.")
    add_bullet_p(doc, "Minggu 11: ", "Uji coba lapangan (UAT) di 1 pilot outlet Kopi Jodi.")
    add_bullet_p(doc, "Minggu 12: ", "Pelatihan staf (Training) dan Go-Live resmi sistem.")
    
    add_body_p(doc, "Ketentuan Pembayaran:", bold_prefix="Ringkasan Ketentuan: ")
    add_bullet_p(doc, "Skema Beli Putus (Turnkey): ", "DP 30% saat SPK, 40% setelah modul siap UAT di pilot outlet, dan pelunasan 30% setelah training & go-live.")
    add_bullet_p(doc, "Skema Bulanan (Dedicated Dev): ", "Rp 5.000.000,- dibayarkan di awal bulan (in-advance) dengan masa kontrak minimal 5 tahun (60 bulan).")

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------
    # SECTION 6: LEMBAR PERSETUJUAN
    # -------------------------------------------------------------
    style_heading(doc.add_paragraph(), "6. Lembar Konfirmasi & Persetujuan", level=1)
    add_body_p(doc, "Silakan beri tanda centang pada opsi yang dipilih oleh Manajemen Kopi Jodi:")

    opt_table = doc.add_table(rows=4, cols=3)
    opt_headers = ["Pilihan", "Opsi Skema Investasi", "Keterangan"]
    opt_widths = [0.8, 3.4, 2.07]
    opt_align = [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT]
    
    for i, h in enumerate(opt_headers):
        opt_table.cell(0, i).paragraphs[0].text = h
        
    opt_data = [
        ("[   ]", "OPSI 1: Beli Putus - Paket Bundling (Fase 1 + 2 + 3)\nInvestasi: Rp 135.000.000,- (Hemat Rp 15 Juta)", "Garansi 6 Bulan, penyerahan source code, sistem tuntas."),
        ("[   ]", "OPSI 1: Beli Putus - Fase 1 Saja (Core Operasional)\nInvestasi: Rp 75.000.000,-", "Fokus operasional outlet, stok otomatis, & kasir live."),
        ("[   ]", "OPSI 2: Sistem Bulanan Dedicated Programmer\nBiaya: Rp 5.000.000,- / bulan", "Min. Kontrak 5 Tahun (60 Bulan).\nBebas tambah fitur apapun & dev berkelanjutan oleh genossys.")
    ]
    
    for r_idx, d in enumerate(opt_data, start=1):
        row = opt_table.rows[r_idx]
        for c_idx, val in enumerate(d):
            row.cells[c_idx].paragraphs[0].text = val
            
    format_table(opt_table, opt_widths, opt_align, header_bg="2C1810")
    
    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Signature Table with Genossys
    sig_table = doc.add_table(rows=4, cols=2)
    sig_widths = [3.13, 3.14]
    
    sig_table.cell(0, 0).paragraphs[0].add_run("Pihak Pertama (Klien):").bold = True
    sig_table.cell(0, 0).add_paragraph("Manajemen / Pemilik Bisnis Kopi Jodi")
    
    sig_table.cell(0, 1).paragraphs[0].add_run("Pihak Kedua (Pengembang):").bold = True
    p_dev_sig = sig_table.cell(0, 1).add_paragraph("genossys (Technology Partner)")
    p_dev_sig.add_run("\nEmail: genossys2019@gmail.com")
    
    sig_table.cell(1, 0).add_paragraph("\n\n( ___________________________________ )")
    sig_table.cell(1, 1).add_paragraph("\n\n( ___________________________________ )")
    
    sig_table.cell(2, 0).add_paragraph("Jabatan: Direktur / Owner")
    sig_table.cell(2, 1).add_paragraph("Jabatan: Lead Solution Architect / genossys")
    
    sig_table.cell(3, 0).add_paragraph("Tanggal: _____ / ______________ / 2026")
    sig_table.cell(3, 1).add_paragraph("Tanggal: _____ / ______________ / 2026")
    
    for r in sig_table.rows:
        for c_idx, c in enumerate(r.cells):
            c.width = Inches(sig_widths[c_idx])
            set_cell_background(c, "FFFFFF")
            set_cell_margins(c, top=50, bottom=50, left=60, right=60)
            set_cell_borders(c, top="E2D9CC", bottom="E2D9CC", left="none", right="none", sz="4")
            for p in c.paragraphs:
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                for r_run in p.runs:
                    r_run.font.name = 'Segoe UI'
                    r_run.font.size = Pt(9)
                    
    try:
        doc.save(file_path)
        print(f"Document successfully created at: {file_path}")
    except PermissionError:
        alt_path = file_path.replace(".docx", "_Terbaru.docx")
        doc.save(alt_path)
        print(f"File asli sedang dibuka di Microsoft Word. Dokumen terbaru disimpan di: {alt_path}")

if __name__ == "__main__":
    out_path = r"c:\PROJECT\WEBSITE\kopi-jodi\Penawaran_Pengembangan_Ekosistem_Aplikasi_Kopi_Jodi.docx"
    build_proposal_docx(out_path)
